import csv
import openpyxl
from docx import Document as DocxDocument
from pypdf import PdfReader


def _parse_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def _parse_xlsx(file_path: str) -> str:
    workbook = openpyxl.load_workbook(
        file_path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append(",".join(str(cell)
                         for cell in row if cell is not None))
    return "\n".join(lines)


def _parse_csv(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        return "\n".join(",".join(row) for row in reader)


def _parse_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


# 按扩展名分发到对应的解析函数，新增一种格式只需要往这个字典里加一行，不用改调用方的代码，这个设计思路和 Day 5 里 _TOOLS_BY_NAME 是一致的。
_PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "csv": _parse_csv,
    "md": _parse_text,
    "txt": _parse_text,
}


def parse_document(file_path: str, file_type: str) -> str:
    """
    解析文档并返回完整文本，不支持的格式或者解析过程中出现异常都会抛出异常，由调用方决定如何处理。
    """
    parser = _PARSERS.get(file_type.lower())
    if parser is None:
        raise ValueError(f"不支持的文档格式：{file_type}")
    return parser(file_path)
